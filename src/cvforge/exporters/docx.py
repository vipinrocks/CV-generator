"""DOCX exporter using python-docx.

Produces a clean, editable document with proper heading styles and bullet
lists so recruiters and ATS software can parse it easily.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from cvforge.markdown_lite import TextRun, to_runs
from cvforge.models import CV, CVForgeError, ExperienceItem
from cvforge.photo import PhotoError, circular_photo_data_uri

ACCENT = RGBColor(0x25, 0x63, 0xEB)
MUTED = RGBColor(0x55, 0x55, 0x55)


class DocxExportError(CVForgeError):
    """Raised when the DOCX file cannot be written."""


def _set_style(doc: DocumentObject, name: str, size: int | None = None, **kwargs: Any) -> None:
    """Adjust a built-in style's font, size and colour if it exists."""
    try:
        style = doc.styles[name]
    except KeyError:
        return
    font = style.font
    if size is not None:
        font.size = Pt(size)
    color = kwargs.get("color")
    if color is not None:
        font.color.rgb = color
    bold = kwargs.get("bold")
    if bold is not None:
        font.bold = bold
    name_font = kwargs.get("font_name")
    if name_font:
        font.name = name_font


def _add_runs(paragraph: Any, text: str) -> None:
    """Add markdown-lite formatted runs to a paragraph, resolving hyperlinks."""
    for run in to_runs(text):
        _add_run(paragraph, run)


def _add_run(paragraph: Any, run: TextRun) -> None:
    """Add a single styled run, as a hyperlink if it carries a URL."""
    if run.link:
        run_element = _add_hyperlink(paragraph, run.text, run.link)
        styler = _RunStyle(run_element)
        if run.bold:
            styler.bold = True
        if run.italic:
            styler.italic = True
        return
    r = paragraph.add_run(run.text)
    r.bold = run.bold or None
    r.italic = run.italic or None


def _add_hyperlink(paragraph: Any, text: str, url: str) -> Any:
    """Append a clickable hyperlink to *paragraph* and return its ``w:r`` element."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    r_pr.append(r_style)
    new_run.append(r_pr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return new_run


class _RunStyle:
    """Sets bold/italic directly on an lxml ``w:r`` element."""

    def __init__(self, element: Any) -> None:
        self._element = element

    def _set_prop(self, prop: str, value: bool) -> None:
        from docx.oxml import OxmlElement

        r_pr = self._element.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            self._element.insert(0, r_pr)
        tag = qn(f"w:{prop}")
        existing = r_pr.find(tag)
        if existing is None:
            existing = OxmlElement(f"w:{prop}")
            r_pr.append(existing)
        existing.set(qn("w:val"), "1" if value else "0")

    @property
    def bold(self) -> None:
        return None

    @bold.setter
    def bold(self, value: bool | None) -> None:
        if value:
            self._set_prop("b", True)

    @property
    def italic(self) -> None:
        return None

    @italic.setter
    def italic(self, value: bool | None) -> None:
        if value:
            self._set_prop("i", True)


def _add_section_heading(doc: DocumentObject, title: str) -> None:
    """Add a section heading paragraph."""
    doc.add_heading(title, level=1)


def _add_experience(doc: DocumentObject, items: list[ExperienceItem]) -> None:
    """Add experience entries with role/company lines and bullet highlights."""
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(f"{item.role} — {item.company}")
        run.bold = True
        meta = []
        if item.start_date:
            meta.append(f"{item.start_date} – {item.end_date or 'Present'}")
        if item.location:
            meta.append(item.location)
        if meta:
            p_meta = doc.add_paragraph()
            p_meta.paragraph_format.space_after = Pt(2)
            r = p_meta.add_run(" · ".join(meta))
            r.italic = True
            r.font.color.rgb = MUTED
            p_meta.paragraph_format.left_indent = Pt(12)
        for highlight in item.highlights:
            bullet = doc.add_paragraph(style="List Bullet")
            bullet.paragraph_format.space_after = Pt(0)
            _add_runs(bullet, highlight)


def export_docx(
    cv: CV,
    out_path: str | Path,
    theme: str | None = None,
    options: Any | None = None,
) -> Path:
    """Export *cv* to an editable DOCX file.

    Args:
        cv: The validated CV data.
        out_path: Destination file path (``.docx``).
        theme: Ignored; present so DOCX/HTML/PDF exporters share a signature.
        options: Ignored for the same reason.

    Returns:
        The resolved output path.

    Raises:
        DocxExportError: If the file cannot be written.
    """
    doc = Document()

    # Base typography: clean and ATS-friendly.
    _set_style(doc, "Normal", size=10, font_name="Calibri")
    _set_style(doc, "Heading 1", size=13, bold=True, color=ACCENT)
    _set_style(doc, "Heading 2", size=11, bold=True, color=ACCENT)
    _set_style(doc, "Title", size=20, bold=True)

    basics = cv.basics
    name_p = doc.add_paragraph()
    name_run = name_p.add_run(basics.full_name)
    name_run.bold = True
    name_run.font.size = Pt(22)
    if basics.headline:
        head_p = doc.add_paragraph()
        head_run = head_p.add_run(basics.headline)
        head_run.italic = True
        head_run.font.size = Pt(12)

    contact_bits = [b for b in (basics.email, basics.phone, basics.location) if b]
    links = [b for b in (basics.website, basics.linkedin, basics.github) if b]
    if contact_bits or links:
        contact_p = doc.add_paragraph()
        contact_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for i, bit in enumerate(contact_bits + links):
            if i:
                contact_p.add_run(" · ")
            if bit in links:
                _add_hyperlink(contact_p, bit, bit)
            else:
                contact_p.add_run(bit)
        for run in contact_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = MUTED

    if basics.photo:
        try:
            uri = circular_photo_data_uri(basics.photo)
            _add_data_uri_image(doc, uri)
        except PhotoError:
            pass  # a broken photo must never block the export

    if basics.summary:
        _add_section_heading(doc, "Summary")
        summary_p = doc.add_paragraph()
        _add_runs(summary_p, basics.summary)

    if cv.experience:
        _add_section_heading(doc, "Experience")
        _add_experience(doc, cv.experience)

    if cv.education:
        _add_section_heading(doc, "Education")
        for ed in cv.education:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(ed.degree + (f", {ed.field}" if ed.field else ""))
            run.bold = True
            sub = doc.add_paragraph()
            sub.paragraph_format.space_after = Pt(2)
            bits = [ed.institution]
            if ed.start_date:
                bits.append(f"{ed.start_date} – {ed.end_date or ''}".rstrip(" –"))
            if ed.grade:
                bits.append(ed.grade)
            r = sub.add_run(" · ".join(bits))
            r.italic = True
            r.font.color.rgb = MUTED
            if ed.notes:
                _add_runs(doc.add_paragraph(), ed.notes)

    if cv.skills:
        _add_section_heading(doc, "Skills")
        for group in cv.skills:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            label = p.add_run(f"{group.name}: ")
            label.bold = True
            p.add_run(", ".join(group.skills))

    if cv.projects:
        _add_section_heading(doc, "Projects")
        for proj in cv.projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            if proj.url:
                _add_hyperlink(p, proj.name, proj.url)
            else:
                run = p.add_run(proj.name)
                run.bold = True
            if proj.tech:
                r = p.add_run(f"  ({', '.join(proj.tech)})")
                r.italic = True
                r.font.color.rgb = MUTED
            if proj.description:
                _add_runs(doc.add_paragraph(), proj.description)
            for highlight in proj.highlights:
                bullet = doc.add_paragraph(style="List Bullet")
                bullet.paragraph_format.space_after = Pt(0)
                _add_runs(bullet, highlight)

    if cv.certifications:
        _add_section_heading(doc, "Certifications")
        for cert in cv.certifications:
            p = doc.add_paragraph(style="List Bullet")
            if cert.url:
                _add_hyperlink(p, cert.name, cert.url)
            else:
                p.add_run(cert.name)
            bits = [b for b in (cert.issuer, cert.date) if b]
            if bits:
                r = p.add_run(f" — {' · '.join(bits)}")
                r.font.color.rgb = MUTED

    if cv.languages:
        _add_section_heading(doc, "Languages")
        for lang in cv.languages:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(lang.language + (f" — {lang.proficiency}" if lang.proficiency else ""))

    if cv.awards:
        _add_section_heading(doc, "Awards")
        for award in cv.awards:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(award.name)
            run.bold = True
            bits = [b for b in (award.issuer, award.date) if b]
            if bits:
                p.add_run(f" — {' · '.join(bits)}")
            if award.description:
                _add_runs(doc.add_paragraph(), award.description)

    if cv.publications:
        _add_section_heading(doc, "Publications")
        for pub in cv.publications:
            p = doc.add_paragraph(style="List Bullet")
            if pub.url:
                _add_hyperlink(p, pub.title, pub.url)
            else:
                p.add_run(pub.title)
            bits = [b for b in (pub.venue, pub.date) if b]
            if bits:
                r = p.add_run(f" — {' · '.join(bits)}")
                r.font.color.rgb = MUTED

    if cv.volunteering:
        _add_section_heading(doc, "Volunteering")
        for vol in cv.volunteering:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(vol.role or vol.organization)
            run.bold = True
            meta = [b for b in (vol.organization if vol.role else "", vol.start_date) if b]
            if meta:
                r = p.add_run(f" — {' · '.join(meta)}")
                r.font.color.rgb = MUTED
            for highlight in vol.highlights:
                bullet = doc.add_paragraph(style="List Bullet")
                _add_runs(bullet, highlight)

    if cv.interests:
        _add_section_heading(doc, "Interests")
        p = doc.add_paragraph()
        p.add_run(" · ".join(cv.interests))

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out))
    except OSError as exc:
        raise DocxExportError(f"Could not write '{out}': {exc}") from exc
    return out


def _add_data_uri_image(doc: DocumentObject, data_uri: str) -> None:
    """Decode a base64 data URI and insert it as a centred picture."""
    import base64
    import io

    from PIL import Image

    header, _, b64 = data_uri.partition(",")
    if not b64:
        return
    fmt = "PNG" if "png" in header else "JPEG"
    image = Image.open(io.BytesIO(base64.b64decode(b64)))
    buffer = io.BytesIO()
    image.convert("RGB").save(buffer, format=fmt)
    buffer.seek(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(buffer, width=Pt(72))
